"""FAQ 知识类型（客服改造第 5 项）：标准问 + 答案 + 相似问数组。

运营维护 FAQ（Excel 导入），客服提问先 FAQ 向量匹配（阈值内直答），
未命中再走文档 RAG——FAQ 直答快且准，是客服高频问题的第一道命门。

存储：SQLite（持久化）+ 内存向量索引（启动加载，FAQ 量级小）。
向量：复用 EmbeddingClient，question + 相似问都嵌入，取最大相似度。
"""
from __future__ import annotations

import json
import logging
import sqlite3
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def _cosine(a: list[float], b: list[float]) -> float:
    """余弦相似度。向量维度不一致返回 0。"""
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    na = sum(x * x for x in a) ** 0.5
    nb = sum(x * x for x in b) ** 0.5
    if na == 0 or nb == 0:
        return 0.0
    return dot / (na * nb)


class FAQStore:
    """FAQ 存储与向量匹配。"""

    def __init__(self, db_path: str | Path, embeddings: Any):
        self._embeddings = embeddings
        self._conn = sqlite3.connect(str(db_path), check_same_thread=False)
        self._conn.execute("PRAGMA journal_mode=WAL")
        self._conn.execute("PRAGMA busy_timeout=5000")
        self._conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS faq (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                kb_id       TEXT NOT NULL DEFAULT 'default',
                question    TEXT NOT NULL,
                answer      TEXT NOT NULL,
                similar     TEXT NOT NULL DEFAULT '[]',
                created_at  TEXT DEFAULT (datetime('now'))
            );
            """
        )
        self._conn.commit()
        # 内存向量索引：kb_id -> [(faq_id, question, answer, vector), ...]
        self._index: dict[str, list[tuple[int, str, str, list[float]]]] = {}
        self._load_index()

    def _load_index(self) -> None:
        """启动时把 FAQ 加载进内存向量索引。"""
        rows = self._conn.execute(
            "SELECT id, kb_id, question, answer, similar FROM faq"
        ).fetchall()
        for fid, kb_id, q, a, sim in rows:
            try:
                similars = json.loads(sim or "[]")
            except Exception:
                similars = []
            texts = [q] + [s for s in similars if s]
            try:
                vectors = self._embeddings.embed_texts(texts)
            except Exception as exc:
                logger.warning("FAQ 向量化失败，跳过 id=%s: %s", fid, exc)
                continue
            # 每个文本一个向量条目（question + 相似问）
            for text, vec in zip(texts, vectors):
                self._index.setdefault(kb_id, []).append((fid, text, a, vec))
        logger.info("FAQ 索引加载：%d 条", len(rows))

    def add_faq(
        self,
        kb_id: str,
        question: str,
        answer: str,
        similar_questions: list[str] | None = None,
    ) -> int:
        """新增 FAQ。返回 faq_id。"""
        similars = [s for s in (similar_questions or []) if s.strip()]
        cur = self._conn.execute(
            "INSERT INTO faq(kb_id, question, answer, similar) VALUES(?,?,?,?)",
            (kb_id, question, answer, json.dumps(similars, ensure_ascii=False)),
        )
        self._conn.commit()
        fid = cur.lastrowid
        # 更新内存索引
        texts = [question] + similars
        try:
            vectors = self._embeddings.embed_texts(texts)
            for text, vec in zip(texts, vectors):
                self._index.setdefault(kb_id, []).append((fid, text, answer, vec))
        except Exception as exc:
            logger.warning("FAQ 向量化失败（已入库，索引未更新）: %s", exc)
        return fid

    def search(self, query: str, kb_id: str, threshold: float = 0.8) -> dict[str, Any] | None:
        """FAQ 向量匹配。阈值内返回 {question, answer, score}，否则 None。"""
        entries = self._index.get(kb_id, [])
        if not entries:
            return None
        try:
            qvec = self._embeddings.embed_query(query)
        except Exception as exc:
            logger.warning("FAQ 查询向量化失败: %s", exc)
            return None
        best: tuple[float, int, str, str] | None = None
        for fid, text, answer, vec in entries:
            score = _cosine(qvec, vec)
            if best is None or score > best[0]:
                best = (score, fid, text, answer)
        if best is None or best[0] < threshold:
            return None
        return {"question": best[2], "answer": best[3], "score": round(best[0], 4)}

    def list_faqs(self, kb_id: str | None = None) -> list[dict[str, Any]]:
        """列出 FAQ（管理界面/Excel 导出用）。"""
        if kb_id:
            rows = self._conn.execute(
                "SELECT id, kb_id, question, answer, similar FROM faq WHERE kb_id=? ORDER BY id",
                (kb_id,),
            ).fetchall()
        else:
            rows = self._conn.execute(
                "SELECT id, kb_id, question, answer, similar FROM faq ORDER BY id"
            ).fetchall()
        result = []
        for fid, k, q, a, sim in rows:
            try:
                similars = json.loads(sim or "[]")
            except Exception:
                similars = []
            result.append(
                {
                    "id": fid,
                    "kb_id": k,
                    "question": q,
                    "answer": a,
                    "similar_questions": similars,
                }
            )
        return result

    def delete_faq(self, faq_id: int) -> int:
        """删除 FAQ。返回删除条数。

        P2：只移除该 faq 的索引条目，不全量重建（避免删除一次就全库重 embedding）。
        """
        row = self._conn.execute("SELECT kb_id FROM faq WHERE id=?", (faq_id,)).fetchone()
        cur = self._conn.execute("DELETE FROM faq WHERE id=?", (faq_id,))
        self._conn.commit()
        if cur.rowcount > 0 and row:
            kb_id = row[0]
            self._index[kb_id] = [
                e for e in self._index.get(kb_id, []) if e[0] != faq_id
            ]
        return cur.rowcount

    # ---------- Excel 批量导入导出（运营人员维护）----------
    @staticmethod
    def _split_similars(raw: str) -> list[str]:
        """相似问拆分：分号/逗号/换行分隔，去空白去空。"""
        import re

        return [s.strip() for s in re.split(r"[;；,，\n]", raw or "") if s.strip()]

    def import_excel(self, file_bytes: bytes, kb_id: str) -> dict[str, Any]:
        """从 Excel 批量导入 FAQ。模板：标准问 | 答案 | 相似问（分号分隔）。

        行级校验：标准问、答案必填；跳过表头；逐行导入，单行失败不影响其它行。
        返回 {imported, skipped, errors: [{row, reason}]}。
        """
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
        ws = wb.active
        imported, skipped = 0, 0
        errors: list[dict[str, Any]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1:  # 表头
                continue
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue  # 空行
            question = str(row[0]).strip() if row[0] is not None else ""
            answer = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            similars_raw = str(row[2]) if len(row) > 2 and row[2] is not None else ""
            if not question or not answer:
                skipped += 1
                errors.append({"row": i, "reason": "标准问或答案为空"})
                continue
            try:
                self.add_faq(kb_id, question, answer, self._split_similars(similars_raw))
                imported += 1
            except Exception as exc:
                skipped += 1
                errors.append({"row": i, "reason": str(exc)})
        return {"imported": imported, "skipped": skipped, "errors": errors}

    def export_excel(self, kb_id: str | None = None) -> bytes:
        """导出 FAQ 为 Excel 字节流（模板同导入）。"""
        import io

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "FAQ"
        ws.append(["标准问", "答案", "相似问（分号分隔）"])
        for faq in self.list_faqs(kb_id=kb_id):
            ws.append(
                [faq["question"], faq["answer"], "；".join(faq["similar_questions"])]
            )
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
