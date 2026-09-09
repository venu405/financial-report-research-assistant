"""KB 入库链路测试：分块 + 元数据（含 kb_id）+ 多类型解析（md/txt/pdf/docx/图片/扫描PDF）。"""
from __future__ import annotations

import os

import pytest

from services.kb.ingest import (
    build_chunks,
    chunk_text,
    parse_document,
    parse_document_structured,
)

# ---- OCR 相关辅助：依赖/模型缺失时跳过（离线/CI 不硬失败；本机已预热模型缓存则正常跑） ----


@pytest.fixture(scope="module")
def ocr_ready():
    """OCR 引擎可初始化才继续；否则跳过（缺包或缺本地模型）。"""
    from services.kb.ocr import get_ocr_engine

    try:
        get_ocr_engine()
    except ValueError as exc:
        pytest.skip(f"OCR 不可用，跳过：{exc}")
    return True


_CJK_FONTS = [
    # Windows
    r"C:/Windows/Fonts/msyh.ttc",
    r"C:/Windows/Fonts/simhei.ttf",
    r"C:/Windows/Fonts/simsun.ttc",
    # Linux（CI/Docker，装 fonts-noto-cjk 后可用）
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    "/usr/share/fonts/wqy-zenhei/wqy-zenhei.ttc",
    # macOS
    "/System/Library/Fonts/PingFang.ttc",
]


def _find_cjk_font() -> str | None:
    return next((p for p in _CJK_FONTS if os.path.exists(p)), None)


def _make_text_image(path, text: str, font_size: int = 48) -> None:
    """用 Pillow + 系统中文字体，把文字渲染成一张白底黑字图片。"""
    from PIL import Image, ImageDraw, ImageFont

    font_path = _find_cjk_font()
    if not font_path:
        pytest.skip("系统缺少中文字体，无法生成 OCR 测试图")
    img = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(img).text(
        (30, 70), text, font=ImageFont.truetype(font_path, font_size), fill="black"
    )
    img.save(str(path))


def test_chunk_text_overlap():
    text = "一" * 2000
    chunks = chunk_text(text, chunk_size=800, overlap=100)
    assert len(chunks) >= 3  # 2000 字符按 700 步进至少 3 块
    assert all(len(c) <= 800 for c in chunks)
    # 相邻块有重叠（防止语义在边界被切断）
    assert chunks[0][-100:] == chunks[1][:100]


def test_build_chunks_metadata(tmp_path):
    f = tmp_path / "制度.md"
    f.write_text("# 采购制度\n超过5万元必须招投标。\n", encoding="utf-8")
    chunks = build_chunks(f, doc_id="abc123", kb_id="hr")
    assert chunks
    c = chunks[0]
    assert c.doc_id == "abc123"
    assert c.kb_id == "hr"
    assert c.source_type == "md"
    assert c.chunk_index == 0


def test_legacy_chunk_profile_keeps_rollback_shape(tmp_path):
    from services.kb.ingest import build_legacy_chunks

    path = tmp_path / "旧版.md"
    path.write_text("# 第一章\n" + "旧版正文" * 80, encoding="utf-8")

    chunks = build_legacy_chunks(path, doc_id="legacy", chunk_size=120, overlap=20)

    assert len(chunks) > 1
    assert all(chunk.page_start is None for chunk in chunks)
    assert all(chunk.section_path == "" for chunk in chunks)
    assert all(chunk.metadata == {} for chunk in chunks)


def test_parse_markdown(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("# 标题\n正文内容", encoding="utf-8")
    assert "标题" in parse_document(f)


def test_structured_markdown_keeps_section_path_and_legacy_text(tmp_path):
    f = tmp_path / "章节.md"
    f.write_text("# 采购制度\n总则。\n\n## 审批流程\n五万元以上须审批。", encoding="utf-8")

    blocks = parse_document_structured(f)

    assert isinstance(parse_document(f), str)
    assert blocks[0].section_path == "采购制度"
    assert blocks[1].section_path == "采购制度 > 审批流程"
    assert all(block.page_start is None for block in blocks)


def test_build_chunks_parent_child_and_neighbor_metadata(tmp_path):
    f = tmp_path / "父子.md"
    f.write_text(
        "# 采购制度\n\n第一段说明采购范围和审批规则。\n\n"
        "第二段说明采购金额和审批人。\n\n第三段说明例外情况和留痕要求。",
        encoding="utf-8",
    )

    chunks = build_chunks(f, doc_id="pc1", parent_chunk_size=80, child_chunk_size=28)
    parents = [chunk for chunk in chunks if chunk.chunk_type == "parent"]
    children = [chunk for chunk in chunks if chunk.chunk_type == "child"]

    assert parents and children
    assert all(child.parent_id for child in children)
    assert all(any(parent.metadata["chunk_id"] == child.parent_id for parent in parents) for child in children)
    assert all(chunk.section_path == "采购制度" for chunk in chunks)
    assert all("chunk_type" in chunk.metadata for chunk in chunks)
    assert children[0].previous_chunk_id == ""
    assert children[-1].next_chunk_id == ""
    assert all("page_start" not in child.metadata for child in children)


def test_markdown_table_header_is_kept_when_split():
    table = "| 字段 | 说明 |\n| --- | --- |\n" + "\n".join(
        f"| 字段{i} | 这是很长的说明{i} |" for i in range(8)
    )

    chunks = chunk_text(table, chunk_size=45, overlap=5)

    assert len(chunks) > 1
    assert all("| 字段 | 说明 |" in chunk for chunk in chunks)


# ---- 多类型解析：图片 OCR / 扫描 PDF / 文字 PDF / docx / 不支持类型 ----


def test_parse_image_ocr(tmp_path, ocr_ready):
    """图片 → OCR 识别成文字（新增能力：图片入库）。"""
    p = tmp_path / "合同扫描.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")
    text = parse_document(p)
    assert "五万元" in text


def test_parse_image_metadata(tmp_path, ocr_ready):
    """图片入库的分块元数据：source_type = 图片扩展名。"""
    p = tmp_path / "制度.png"
    _make_text_image(p, "员工年假每年十五天")
    chunks = build_chunks(p, doc_id="img1")
    assert chunks
    assert chunks[0].source_type == "png"
    assert chunks[0].doc_id == "img1"


def test_parse_scanned_pdf(tmp_path, ocr_ready):
    """扫描版 PDF（无文字层）→ 逐页渲染成图片再 OCR（新增能力）。"""
    import pymupdf
    from pypdf import PdfReader

    img_path = tmp_path / "scan.png"
    _make_text_image(img_path, "项目启动需要总经理报批")
    pdf_path = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(40, 40, 555, 180), filename=str(img_path))
    doc.save(str(pdf_path))
    doc.close()
    # 先确认它确实是"扫描版"：pypdf 提不出文字层
    assert not (PdfReader(str(pdf_path)).pages[0].extract_text() or "").strip()
    chunks = build_chunks(pdf_path, doc_id="scan1")
    assert chunks
    assert chunks[0].source_type == "pdf"
    assert "报批" in chunks[0].text


def test_parse_text_pdf(tmp_path, ocr_ready):
    """文字版 PDF → 提取出文字（pypdf 优先，OCR 兜底）。"""
    import pymupdf

    font_path = _find_cjk_font()
    if not font_path:
        pytest.skip("系统缺少中文字体，无法生成 PDF 测试")
    pdf_path = tmp_path / "text.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    # 嵌入系统 TTF：pypdf 才能干净提取（内置 CID 字体 china-s 提取会乱码）
    page.insert_font(fontname="cjk", fontfile=font_path)
    page.insert_text((72, 72), "这是一段PDF测试文字", fontname="cjk", fontsize=20)
    doc.save(str(pdf_path))
    doc.close()
    text = parse_document(pdf_path)
    assert "PDF测试" in text


def test_pdf_structured_blocks_use_one_based_page_ranges(tmp_path):
    """文字版 PDF 的结构块保留 1 起始页码，并过滤底部孤立页码。"""
    import pymupdf

    pdf_path = tmp_path / "pages.pdf"
    doc = pymupdf.open()
    for index, body in enumerate(("第一章 采购范围", "第二章 审批流程"), start=1):
        page = doc.new_page(width=400, height=300)
        page.insert_text((40, 60), body, fontsize=14)
        page.insert_text((195, 285), str(index), fontsize=10)
    doc.save(str(pdf_path))
    doc.close()

    blocks = parse_document_structured(pdf_path)
    chunks = build_chunks(pdf_path, doc_id="pages1", chunk_size=100)

    assert [block.page_start for block in blocks] == [1, 2]
    assert [chunk.page_start for chunk in chunks] == [1]
    assert [chunk.page_end for chunk in chunks] == [2]
    assert all(chunk.page == chunk.page_start for chunk in chunks)
    assert all(chunk.text.strip() not in {"1", "2"} for chunk in chunks)


def test_parse_docx_paragraph_and_table(tmp_path):
    """docx：段落 + 表格都被提取（既有能力，补测试）。"""
    import docx

    p = tmp_path / "制度.docx"
    d = docx.Document()
    d.add_paragraph("考勤制度第一条：按时上下班")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "部门"
    t.rows[0].cells[1].text = "人数"
    t.rows[1].cells[0].text = "研发部"
    t.rows[1].cells[1].text = "三十人"
    d.save(str(p))

    text = parse_document(p)
    assert "考勤制度" in text  # 段落
    assert "研发部" in text  # 表格


def test_unsupported_extension_raises(tmp_path):
    """不支持的扩展名仍抛 ValueError（如 xlsx）。"""
    p = tmp_path / "a.xlsx"
    p.write_bytes(b"fake")
    with pytest.raises(ValueError, match="不支持"):
        parse_document(p)


# ---- OCR 模式（local/baidu/auto）：mock 后备通道，不触发真实网络 ----


def test_ocr_local_ok_no_fallback(tmp_path, ocr_ready, monkeypatch):
    """auto 模式 + 干净印刷图：本地直接出结果，不调用百度。"""
    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "ok.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")

    called = []
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: called.append(True) or "百度")
    text = ocr_image(p, mode="auto")
    assert "五万元" in text
    assert called == [], "本地结果达标时不应触发百度后备"


def test_ocr_empty_local_uses_baidu(tmp_path, ocr_ready, monkeypatch):
    """auto 模式 + 纯白图（本地识别为空）：走百度后备并返回百度结果。"""
    from PIL import Image

    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "blank.png"
    Image.new("RGB", (300, 100), "white").save(str(p))
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "百度识别的手写内容")
    assert ocr_image(p, mode="auto") == "百度识别的手写内容"


def test_ocr_baidu_unavailable_keeps_local(tmp_path, ocr_ready, monkeypatch):
    """auto 模式 + 本地低置信但百度没帮上忙（返回空）：静默回退本地结果，不抛异常。"""
    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "low.png"
    _make_text_image(p, "潦草字样的测试文字")
    # 强制把阈值抬到 1.0，让正常印刷图也被判"效果差" → 触发后备
    monkeypatch.setattr(ocr, "_fallback_threshold", lambda: 1.0)
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "")
    text = ocr_image(p, mode="auto")
    assert text  # 本地结果被保留
    assert "测试文字" in text


def test_ocr_baidu_error_degrades(tmp_path, ocr_ready, monkeypatch):
    """auto 模式 + 百度后备自身抛异常：静默降级回本地，不炸。"""
    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "err.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")

    def boom(_img):
        raise RuntimeError("baidu module broke")

    monkeypatch.setattr(ocr, "_fallback_threshold", lambda: 1.0)
    monkeypatch.setattr(ocr, "_fallback_baidu", boom)
    assert "五万元" in ocr_image(p, mode="auto")


def test_ocr_baidu_empty_keeps_empty(tmp_path, ocr_ready, monkeypatch):
    """auto 模式 + 本地空 + 百度空：返回空串（不抛异常）。"""
    from PIL import Image

    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "blank2.png"
    Image.new("RGB", (300, 100), "white").save(str(p))
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "")
    assert ocr_image(p, mode="auto") == ""


def test_ocr_mode_local_never_calls_baidu(tmp_path, ocr_ready, monkeypatch):
    """mode="local"（默认）：即使配了百度也不调用；空图返回空串。"""
    from PIL import Image

    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "blank_local.png"
    Image.new("RGB", (300, 100), "white").save(str(p))
    called = []
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: called.append(True) or "百度")
    assert ocr_image(p, mode="local") == ""
    assert called == [], "local 模式绝不调用百度"


def test_ocr_mode_baidu_uses_baidu(tmp_path, ocr_ready, monkeypatch):
    """mode="baidu"：百度结果优先返回。"""
    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "bd.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "百度手写结果")
    assert ocr_image(p, mode="baidu") == "百度手写结果"


def test_ocr_mode_baidu_empty_falls_back_local(tmp_path, ocr_ready, monkeypatch):
    """mode="baidu" 但百度空/未配置：回退本地结果。"""
    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "bd2.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "")
    text = ocr_image(p, mode="baidu")
    assert "五万元" in text


def test_build_chunks_passes_ocr_mode(tmp_path, ocr_ready, monkeypatch):
    """ocr_mode 透传到 ocr_image（图片入库链路：build_chunks → parse_document → OCR）。"""
    import services.kb.ocr as ocr
    from services.kb.ingest import build_chunks

    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "百度文本")
    monkeypatch.setattr(ocr, "_run_local", lambda img: ("", (0.3,)))
    p = tmp_path / "doc.png"
    _make_text_image(p, "采购金额超过五万元必须招投标")
    chunks = build_chunks(p, doc_id="m1", ocr_mode="baidu")
    # mode=baidu → 百度优先 → 直接返回百度结果，验证参数确实传到底层
    assert chunks and chunks[0].text == "百度文本"


def test_result_acceptable_confidence():
    """_result_acceptable 置信度判定：空/低置信不达标，达标/无分数视为可接受。"""
    import services.kb.ocr as ocr

    assert not ocr._result_acceptable("", None)
    assert not ocr._result_acceptable("字", (0.3,))
    assert ocr._result_acceptable("字", (0.9,))
    assert ocr._result_acceptable("字", None)  # 无分数视为可接受


# ---- 内容启发式 _text_looks_bad：潦草触发，正常截图不触发 ----

_SCRAWL_TEXT = """月复一4，一年复一年，确海的
官难、困苦越来越严重，导致国
家有始民不潦生，修窄，在无数
年，元数名革命勇七的努力下，
在有牲做努力有斗下。于19年，
我们的国家建立5，中国，中华
人民共和国成？，。一有多年来"""


def test_text_looks_bad_triggers_on_scrawl():
    """潦草识别文本命中多个异常特征（①汉字夹数字 ②连续标点）→ 判定效果差。"""
    import services.kb.ocr as ocr

    assert ocr._text_looks_bad(_SCRAWL_TEXT)


def test_text_looks_bad_ignores_normal_text():
    """正常文本不触发：含日期/编号/英数/口语双标点都不该被当成乱码。"""
    import services.kb.ocr as ocr

    normal = [
        # 客服截图（印刷体，无异常）
        "我们的机器操作和打印机原理是差不多的您只需要提供要抄写的内容导入就行了。",
        # 诗句（极短但完全正常）
        "愿我如星君如月\n夜夜流光相皎洁",
        # 拼多多截图（英数密集，但不该被当成乱码）
        "S23/s23+/S23PIus/S230ftra\n适用三星S23系列电池\n拼多多全程保障购物体验\n"
        "林梵高，155****0669，广东省深圳市龙岗区\n到手价¥58",
        # 客服长文（口语标点多但无连续异常标点）
        "收到我们的商品以后，售后客服会给您提供试用账号，最长可以试用十五天哦！",
        # 口语强调（连续同款标点是正常表达，不是乱码）
        "好的！！",
        "请问？？",
        # 含日期/编号的正常文档（企业知识库常见）
        "2026年8月17日 采购编号2026-001 需法务审核",
        "第3条 金额5万 采购流程",
        # 中英混写 / 品牌型号
        "iPhone 15 Pro 128GB 黑色",
        "项目编号PROJ-2026-001 型号X100",
        "英语English句子mixed混写",
        # 纯英文
        "Hello world, this is English text 123.",
    ]
    for t in normal:
        assert not ocr._text_looks_bad(t), t


def test_text_looks_bad_empty():
    """空串/纯空白不算乱码。"""
    import services.kb.ocr as ocr

    assert not ocr._text_looks_bad("")
    assert not ocr._text_looks_bad("   \n  ")


def test_ocr_auto_bad_content_uses_baidu(tmp_path, ocr_ready, monkeypatch):
    """auto 模式：置信度达标但文本疑似潦草/乱码 → 仍触发百度兜底。"""
    from PIL import Image

    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "scrawl.png"
    Image.new("RGB", (300, 100), "white").save(str(p))

    # 高置信度（0.9 远超阈值），但内容是潦草错字
    monkeypatch.setattr(ocr, "_run_local", lambda img: (_SCRAWL_TEXT, (0.9,)))
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: "日复一日一年复一年确情的灾难")
    assert ocr_image(p, mode="auto") == "日复一日一年复一年确情的灾难"


def test_ocr_auto_good_content_skips_baidu(tmp_path, ocr_ready, monkeypatch):
    """auto 模式：置信度达标且内容正常 → 直接返回本地结果，不触发百度。"""
    from PIL import Image

    import services.kb.ocr as ocr
    from services.kb.ocr import ocr_image

    p = tmp_path / "good.png"
    Image.new("RGB", (300, 100), "white").save(str(p))

    called = []
    monkeypatch.setattr(ocr, "_run_local", lambda img: ("采购金额超过五万元必须招投标", (0.9,)))
    monkeypatch.setattr(ocr, "_fallback_baidu", lambda img: called.append(True) or "百度")
    assert "五万元" in ocr_image(p, mode="auto")
    assert called == [], "内容正常时不应触发百度后备"


def test_layout_cleaning_removes_only_bottom_page_number():
    """底部孤立页码应清除，正文中的数字和非底部序号必须保留。"""
    from services.kb.ingest import _clean_layout_lines

    lines = [
        (20, 100, "境内外会计准则下会计数据差异"),
        (20, 140, "□适用 √不适用"),
        (20, 200, "营业收入 12 万元"),
        (20, 980, "12"),
    ]
    assert _clean_layout_lines(lines, 1000) == [
        "境内外会计准则下会计数据差异",
        "适用：否",
        "营业收入 12 万元",
    ]


def test_heading_detection_rejects_annual_report_numbers_and_table_cells():
    """年份、金额、电话和 PDF 页码不能被误判为章节，否则真实年报会碎成短块。"""
    from services.kb.ingest import _heading_info

    for text in (
        "2024 年年度报告",
        "2025 年4 月",
        "8 / 300",
        "787,733,705.74",
        "2.19%",
        "0.00",
        "010-53223377",
        "2024年",
        "1、这是从 PDF 正文抽取出来且在句号之前意外换行的一段很长说明文字，不应当被当成新的章节标题从上下文里拆开",
    ):
        assert _heading_info(text) is None

    assert _heading_info("第一章 采购范围") == (1, "第一章 采购范围")
    assert _heading_info("一、主要会计数据") == (1, "一、主要会计数据")
    assert _heading_info("1、报告期主要经营指标") == (1, "1、报告期主要经营指标")
    assert _heading_info("1.2 风险提示") == (2, "1.2 风险提示")


def test_numeric_table_rows_do_not_create_tiny_structured_blocks():
    """连续数字表格行应留在正文块中，不能一行变成一个 section。"""
    from services.kb.ingest import _structured_blocks_from_pages

    rows = ["第一章 主要数据"] + [f"{2024 - index}年 {1000000 + index:,}.00" for index in range(40)]
    blocks = _structured_blocks_from_pages([(1, "\n".join(rows))])

    assert len(blocks) == 1
    assert blocks[0].section_path == "第一章 主要数据"
    assert "1,000,039.00" in blocks[0].text


def test_adjacent_small_sections_are_packed_to_chunk_size_with_page_span():
    """真实年报的小标题不能各占一个几十字 chunk，应聚合并保留跨页范围。"""
    from services.kb.ingest import StructuredBlock, _build_chunk_specs

    blocks = [
        StructuredBlock(
            text=f"{index}、小节标题\n" + "经营数据" * 30,
            page_start=index,
            page_end=index,
            section_path=f"年度报告 > 小节{index}",
        )
        for index in range(1, 7)
    ]

    specs = _build_chunk_specs(
        blocks,
        chunk_size=500,
        overlap=50,
        parent_chunk_size=None,
        child_chunk_size=None,
        include_parent_chunks=False,
    )

    assert len(specs) == 2
    assert specs[0]["page_start"] == 1
    assert specs[0]["page_end"] >= 3
    assert specs[0]["section_path"] == "年度报告"
    assert all(len(spec["text"]) <= 500 for spec in specs)


def test_checkbox_normalization_preserves_decision():
    """勾选行转成明确结论，供向量检索和模型回答使用。"""
    from services.kb.ingest import _normalize_checkbox_text

    assert _normalize_checkbox_text("√是 □否") == "结论：是"
    assert _normalize_checkbox_text("□是 √否") == "结论：否"
    assert _normalize_checkbox_text("□适用 √不适用") == "适用：否"


def test_render_ocr_layout_keeps_table_columns_and_removes_footer():
    """扫描表格按同一行坐标还原列关系，页脚页码不进入结果。"""
    from services.kb.ingest import _render_ocr_layout
    from services.kb.ocr import OCRLine

    lines = [
        OCRLine("事项", 20, 20, 80, 45),
        OCRLine("是否", 500, 20, 550, 45),
        OCRLine("是否适用", 20, 80, 150, 105),
        OCRLine("□适用 √不适用", 500, 80, 700, 105),
        OCRLine("12", 350, 950, 370, 975),
    ]
    text = _render_ocr_layout(lines, 800, 1000)
    assert "事项 | 是否" in text
    assert "是否适用 | 适用：否" in text
    assert "12" not in text
